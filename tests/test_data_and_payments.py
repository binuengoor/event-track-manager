import io
import os
import csv
import json
import zipfile
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.config import settings
from app.services.db_service import db_service


@pytest.fixture
def admin():
    c = TestClient(app)
    c.headers.update({"X-Admin-PIN": settings.admin_pin})
    return c


@pytest.fixture
def client():
    return TestClient(app)


def _signup(client, name, partner=None, song="Song", phone="6105550100", ptype=None):
    perf = {"performance_type": ptype or ("Duet" if partner else "Solo"), "song_title": song}
    if partner:
        perf.update(partner_name=partner, partner_phone="6105550142", partner_age_group="Senior")
    res = client.post("/api/signup", json={"performer_name": name, "contact_info": phone, "age_group": "Senior", "performances": [perf]})
    assert res.status_code == 200, res.text


def _tables(admin):
    return {t["id"]: t for t in admin.get("/api/admin/data").json()["tables"]}


# --------------------------------------------------------------------------- tabs follow the event type
def test_full_event_has_songs_singers_run_sheet_and_potluck(admin):
    assert list(_tables(admin)) == ["songs", "singers", "run_sheet", "potluck"]


def test_potluck_only_event_shows_only_the_potluck_tab(admin):
    admin.post("/api/admin/apply-preset", json={"preset": "potluck"})
    assert list(_tables(admin)) == ["potluck"]


def test_show_without_potluck_has_no_potluck_tab(admin):
    admin.post("/api/admin/apply-preset", json={"preset": "show_no_potluck"})
    assert list(_tables(admin)) == ["songs", "singers", "run_sheet"]


def test_data_endpoints_require_the_pin(client):
    assert client.get("/api/admin/data").status_code == 401
    assert client.get("/api/admin/data/download?format=xlsx").status_code == 401
    assert client.get("/api/admin/payments").status_code == 401
    assert client.put("/api/admin/payments", json={"updates": []}).status_code == 401


def test_opening_the_admin_data_view_refreshes_the_files_in_the_event_folder(admin):
    _signup(admin, "Fresh Singer")
    path = os.path.join(settings.storage.event_dir, "signups.xlsx")
    if os.path.exists(path):
        os.remove(path)
    res = admin.get("/api/admin/data").json()
    assert os.path.getsize(path) > 0 and "signups.xlsx" in res["exported_files"]


# --------------------------------------------------------------------------- singers + payments
def test_singers_are_alphabetical_include_partners_and_exclude_potluck_only_guests(admin, client):
    _signup(client, "Zed Zimmer", song="Z Song")
    _signup(client, "Dev Pillai", partner="Neha Pillai", song="Duet Song")
    _signup(client, "Aarav Menon", song="A Song")
    gid = admin.post("/api/admin/food-groups", json={"name": "Mains"}).json()["group_id"]
    iid = admin.post("/api/admin/food-items", json={"name": "Dal", "group_id": gid}).json()["item_id"]
    client.post("/api/signup", json={"performer_name": "Potluck Pat", "contact_info": "6105550111", "registration_type": "food_only",
                                      "food_signup": {"item_id": iid, "dish_description": "Dal"}})
    names = [s["name"] for s in admin.get("/api/admin/payments").json()["singers"]]
    assert names == ["Aarav Menon", "Dev Pillai", "Neha Pillai", "Zed Zimmer"]
    assert "Potluck Pat" not in names


def test_checking_paid_and_saving_persists_and_feeds_the_spreadsheet(admin, client):
    _signup(client, "Aarav Menon")
    _signup(client, "Dev Pillai", partner="Neha Pillai")
    singers = {s["name"]: s for s in admin.get("/api/admin/payments").json()["singers"]}
    assert not any(s["paid"] for s in singers.values())

    res = admin.put("/api/admin/payments", json={"updates": [
        {"participant_id": singers["Aarav Menon"]["participant_id"], "paid": True},
        {"participant_id": singers["Neha Pillai"]["participant_id"], "paid": True, "note": "Zelle"},
    ]})
    assert res.status_code == 200 and res.json()["paid"] == 2 and res.json()["total"] == 3

    after = {s["name"]: s for s in admin.get("/api/admin/payments").json()["singers"]}
    assert after["Aarav Menon"]["paid"] and after["Aarav Menon"]["paid_at"]
    assert after["Neha Pillai"]["payment_note"] == "Zelle"
    assert not after["Dev Pillai"]["paid"]

    singers_tab = _tables(admin)["singers"]
    paid_col = {r["Name"]: r["Paid"] for r in singers_tab["rows"]}
    assert paid_col == {"Aarav Menon": "Yes", "Dev Pillai": "No", "Neha Pillai": "Yes"}

    # unchecking clears it
    admin.put("/api/admin/payments", json={"updates": [{"participant_id": singers["Aarav Menon"]["participant_id"], "paid": False}]})
    again = {s["name"]: s for s in admin.get("/api/admin/payments").json()["singers"]}
    assert not again["Aarav Menon"]["paid"] and again["Aarav Menon"]["paid_at"] == ""


def test_payment_status_survives_a_rename(admin, client):
    _signup(client, "Old Name")
    pid = admin.get("/api/admin/payments").json()["singers"][0]["participant_id"]
    admin.put("/api/admin/payments", json={"updates": [{"participant_id": pid, "paid": True}]})
    assert client.put("/api/performer/rename", json={"old_name": "Old Name", "new_name": "New Name"}).status_code == 200
    singers = admin.get("/api/admin/payments").json()["singers"]
    assert [(s["name"], s["paid"]) for s in singers] == [("New Name", True)]


def test_payments_are_not_available_for_potluck_only_events(admin):
    admin.post("/api/admin/apply-preset", json={"preset": "potluck"})
    assert admin.get("/api/admin/payments").json()["enabled"] is False
    assert admin.put("/api/admin/payments", json={"updates": []}).status_code == 400


def test_unknown_participant_ids_are_ignored(admin, client):
    _signup(client, "Real Person")
    res = admin.put("/api/admin/payments", json={"updates": [{"participant_id": "pt_nope", "paid": True}]})
    assert res.status_code == 200 and res.json()["paid"] == 0


# --------------------------------------------------------------------------- downloads
def test_download_every_format(admin, client):
    _signup(client, "Dev Pillai", partner="Neha Pillai", song="=1+1")
    from openpyxl import load_workbook

    x = admin.get("/api/admin/data/download?format=xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"] and "attachment" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Songs", "Singers", "Run Sheet", "Potluck"]

    c = admin.get("/api/admin/data/download?format=csv&sheet=singers")
    assert c.status_code == 200 and c.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Name", "Age Group", "Phone"] and [r[0] for r in rows[1:]] == ["Dev Pillai", "Neha Pillai"]

    songs = admin.get("/api/admin/data/download?format=csv&sheet=songs").content.decode("utf-8-sig")
    assert "'=1+1" in songs                      # formula neutralized in the download too

    t = admin.get("/api/admin/data/download?format=tsv&sheet=songs")
    assert "\t" in t.content.decode("utf-8-sig").splitlines()[0]

    j = admin.get("/api/admin/data/download?format=json&sheet=songs").json()
    assert j["tab"] == "songs" and j["count"] == 1 and j["rows"][0]["Performer"] == "Dev Pillai"

    z = admin.get("/api/admin/data/download?format=zip")
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert len(names) == 4 and all(n.endswith(".csv") for n in names)


def test_download_rejects_unknown_format_or_tab(admin):
    assert admin.get("/api/admin/data/download?format=pdf").status_code == 400
    assert admin.get("/api/admin/data/download?format=csv&sheet=nope").status_code == 400


def test_potluck_only_download_has_just_the_potluck_tab(admin):
    admin.post("/api/admin/apply-preset", json={"preset": "potluck"})
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(admin.get("/api/admin/data/download?format=xlsx").content))
    assert wb.sheetnames == ["Potluck"]
