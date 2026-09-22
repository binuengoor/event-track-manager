import io
import logging
from typing import Dict, List, Optional, Tuple, Any

from app.config import settings

logger = logging.getLogger("google-sheet-client")


class GoogleSheetClient:
    """Handles raw Google Sheets API interactions and Google Drive-hosted .xlsx workbooks."""

    def __init__(self, sheets_service, drive_service, is_xlsx: bool = False, mock_mode: bool = False):
        self.sheets = sheets_service
        self.drive = drive_service
        self.is_xlsx = is_xlsx
        self.mock_mode = mock_mode

    def get_sheet_row_mapping(self) -> Dict[str, int]:
        """Returns a mapping from entry_id to 1-based row index in the target Google or Excel Sheet."""
        if self.mock_mode or not settings.google.sheet_id:
            return {}

        tab_name = settings.google.sheet_tab_name or "Song Sign-Up"
        row_map = {}
        try:
            if self.is_xlsx:
                import openpyxl
                content = self.drive.files().get_media(fileId=settings.google.sheet_id, supportsAllDrives=True).execute()
                wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
                ws = wb[tab_name] if tab_name in wb.sheetnames else wb.active
                for idx, r in enumerate(ws.iter_rows(values_only=True)):
                    if idx == 0:
                        continue
                    eid = str(r[0]).strip() if r and len(r) > 0 and r[0] is not None else ""
                    if eid:
                        row_map[eid] = idx + 1
            else:
                res = self.sheets.spreadsheets().values().get(
                    spreadsheetId=settings.google.sheet_id,
                    range=f"'{tab_name}'!A1:A100"
                ).execute()
                rows = res.get("values", [])
                for idx, r in enumerate(rows):
                    if idx == 0:
                        continue
                    eid = str(r[0]).strip() if r and len(r) > 0 else ""
                    if eid:
                        row_map[eid] = idx + 1
        except Exception as ex:
            logger.warning("Failed to fetch sheet row mapping: %s", ex)
        return row_map

    def fetch_raw_sheet_data(self) -> Tuple[List[str], List[List[str]]]:
        """Fetches header row and data rows from Google Sheets or Excel workbook."""
        if self.mock_mode or not settings.google.sheet_id:
            return [], []

        tab_name = settings.google.sheet_tab_name or "Song Sign-Up"
        raw_header: List[str] = []
        raw_rows: List[List[str]] = []

        if self.is_xlsx:
            import openpyxl
            content = self.drive.files().get_media(fileId=settings.google.sheet_id, supportsAllDrives=True).execute()
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb[tab_name] if tab_name in wb.sheetnames else wb.active
            all_rows = []
            for r in ws.iter_rows(values_only=True):
                all_rows.append([str(c).strip() if c is not None else "" for c in r])
            if all_rows:
                raw_header = all_rows[0]
                raw_rows = all_rows[1:]
        else:
            result = self.sheets.spreadsheets().values().get(
                spreadsheetId=settings.google.sheet_id,
                range=f"{tab_name}!A1:Z"
            ).execute()
            all_rows = result.get("values", [])
            if all_rows:
                raw_header = all_rows[0]
                raw_rows = all_rows[1:]

        return raw_header, raw_rows

    def update_track_metadata(
        self,
        sheet_row: int,
        file_id: str,
        duration_str: Optional[str],
        now_iso: str
    ) -> bool:
        """Pushes track status, duration, and file ID to the specified sheet row."""
        if self.mock_mode or not settings.google.sheet_id or sheet_row <= 0:
            return True

        cols = settings.columns
        tab_name = settings.google.sheet_range.split("!")[0] if "!" in settings.google.sheet_range else "Song Sign-Up"

        if self.is_xlsx:
            import openpyxl
            from googleapiclient.http import MediaIoBaseUpload

            content = self.drive.files().get_media(fileId=settings.google.sheet_id, supportsAllDrives=True).execute()
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb[tab_name] if tab_name in wb.sheetnames else wb.active

            ws.cell(row=sheet_row, column=cols.track_status + 1, value="Yes")
            if duration_str:
                ws.cell(row=sheet_row, column=cols.duration + 1, value=duration_str)
            ws.cell(row=sheet_row, column=cols.drive_file_id + 1, value=file_id)
            ws.cell(row=sheet_row, column=cols.last_updated + 1, value=now_iso)

            out_buf = io.BytesIO()
            wb.save(out_buf)
            out_buf.seek(0)

            media = MediaIoBaseUpload(out_buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resumable=True)
            self.drive.files().update(fileId=settings.google.sheet_id, media_body=media, supportsAllDrives=True).execute()
            return True
        else:
            def col_letter(col_idx: int) -> str:
                return chr(ord('A') + col_idx)

            updates = [
                {"range": f"{tab_name}!{col_letter(cols.track_status)}{sheet_row}", "values": [["Yes"]]},
                {"range": f"{tab_name}!{col_letter(cols.drive_file_id)}{sheet_row}", "values": [[file_id]]},
                {"range": f"{tab_name}!{col_letter(cols.last_updated)}{sheet_row}", "values": [[now_iso]]}
            ]
            if duration_str:
                updates.append({"range": f"{tab_name}!{col_letter(cols.duration)}{sheet_row}", "values": [[duration_str]]})

            self.sheets.spreadsheets().values().batchUpdate(
                spreadsheetId=settings.google.sheet_id,
                body={"valueInputOption": "USER_ENTERED", "data": updates}
            ).execute()
            return True

    def update_performance_status(self, sheet_row: int, val_to_write: str) -> bool:
        """Updates performance status in Google Sheet or Excel."""
        if self.mock_mode or not settings.google.sheet_id or sheet_row <= 0:
            return True

        cols = settings.columns
        tab_name = settings.google.sheet_range.split("!")[0] if "!" in settings.google.sheet_range else "Song Sign-Up"

        if self.is_xlsx:
            import openpyxl
            from googleapiclient.http import MediaIoBaseUpload

            content = self.drive.files().get_media(fileId=settings.google.sheet_id, supportsAllDrives=True).execute()
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb[tab_name] if tab_name in wb.sheetnames else wb.active

            ws.cell(row=sheet_row, column=cols.performance_status + 1, value=val_to_write)
            out_buf = io.BytesIO()
            wb.save(out_buf)
            out_buf.seek(0)

            media = MediaIoBaseUpload(out_buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resumable=True)
            self.drive.files().update(fileId=settings.google.sheet_id, media_body=media, supportsAllDrives=True).execute()
            return True
        else:
            def col_letter(col_idx: int) -> str:
                return chr(ord('A') + col_idx)

            cell = f"{tab_name}!{col_letter(cols.performance_status)}{sheet_row}"
            self.sheets.spreadsheets().values().update(
                spreadsheetId=settings.google.sheet_id,
                range=cell,
                valueInputOption="USER_ENTERED",
                body={"values": [[val_to_write]]}
            ).execute()
            return True

    def clear_stale_file_rows(self, row_indices: List[int]) -> None:
        """Resets track status, duration, and file ID for entries whose Drive files were removed."""
        if self.mock_mode or not self.sheets or not row_indices or not settings.google.sheet_id:
            return

        cols = settings.columns
        tab_name = settings.google.sheet_range.split("!")[0] if "!" in settings.google.sheet_range else "Song Sign-Up"
        def col_letter(col_idx: int) -> str:
            return chr(ord('A') + col_idx)

        try:
            sheet_updates = []
            for r_idx in row_indices:
                sheet_updates.append({"range": f"{tab_name}!{col_letter(cols.track_status)}{r_idx}", "values": [["Pending"]]})
                sheet_updates.append({"range": f"{tab_name}!{col_letter(cols.duration)}{r_idx}", "values": [[""]]})
                sheet_updates.append({"range": f"{tab_name}!{col_letter(cols.drive_file_id)}{r_idx}", "values": [[""]]})
            self.sheets.spreadsheets().values().batchUpdate(
                spreadsheetId=settings.google.sheet_id,
                body={"valueInputOption": "USER_ENTERED", "data": sheet_updates}
            ).execute()
        except Exception as ex:
            logger.warning("Could not clear stale file rows in Google Sheet: %s", ex)

    def batch_update_sequence_orders(self, updates: List[Dict[str, Any]]) -> None:
        """Batch updates cells in Google Sheet with new sequence orders."""
        if self.mock_mode or not self.sheets or not updates or not settings.google.sheet_id:
            return

        self.sheets.spreadsheets().values().batchUpdate(
            spreadsheetId=settings.google.sheet_id,
            body={"valueInputOption": "USER_ENTERED", "data": updates}
        ).execute()

    def update_xlsx_sequence_orders(self, row_seq_map: Dict[int, int]) -> None:
        """Updates sequence orders in an Excel sheet stored in Google Drive."""
        if self.mock_mode or not self.drive or not row_seq_map or not settings.google.sheet_id:
            return

        import openpyxl
        from googleapiclient.http import MediaIoBaseUpload

        cols = settings.columns
        tab_name = settings.google.sheet_range.split("!")[0] if "!" in settings.google.sheet_range else "Song Sign-Up"

        content = self.drive.files().get_media(fileId=settings.google.sheet_id, supportsAllDrives=True).execute()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb[tab_name] if tab_name in wb.sheetnames else wb.active

        for row_idx, seq in row_seq_map.items():
            if row_idx > 0:
                ws.cell(row=row_idx, column=cols.sequence_order + 1, value=seq)

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)

        media = MediaIoBaseUpload(out_buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resumable=True)
        self.drive.files().update(fileId=settings.google.sheet_id, media_body=media, supportsAllDrives=True).execute()
