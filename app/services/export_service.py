"""Writes the event's signup workbook/CSVs and database snapshots into the event folder.

signups.xlsx / signups.csv / potluck.csv are regenerated from SQLite (the source of truth) and written
atomically (temp file + rename), so a sync tool such as rclone never uploads a half-written file.
"""
import os
import csv
import glob
import sqlite3
import logging
import tempfile
import threading
from datetime import datetime
from typing import Dict, Any, List

from app.config import settings

logger = logging.getLogger("export-service")

SNAPSHOT_MIN_INTERVAL_SECONDS = 300

PERFORMANCE_COLUMNS = [
    ("Entry ID", "entry_id"), ("Sequence", "sequence_order"), ("Performer", "performer_name"),
    ("Age Group", "age_group"), ("Performance Type", "performance_type"), ("Duet Partner", "partner_name"),
    ("Contact Phone", "contact_info"), ("Guardian", "guardian_name"), ("Guardian Phone", "guardian_phone"),
    ("Partner Phone", "partner_phone"), ("Song", "song_title"), ("Movie/Album", "movie_name"),
    ("Performance Status", "performance_status"), ("Track Status", "track_status"), ("Track File", "drive_file_name"),
    ("Media", "media_type"), ("Duration", "duration"), ("Stage Notes", "stage_notes"),
    ("Signed Up Via", "created_via"), ("Last Updated", "last_updated"),
]
POTLUCK_COLUMNS = [
    ("Group", "group_name"), ("Item", "name"), ("Status", "_status"), ("Signed Up By", "signer_name"),
    ("Phone", "signer_phone"), ("Dish", "dish_description"),
]
PARTICIPANT_COLUMNS = [
    ("Name", "name"), ("Phone", "phone"), ("Age Group", "age_group"),
    ("Guardian", "guardian_name"), ("Guardian Phone", "guardian_phone"),
]


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    # Neutralize spreadsheet formula injection from user-supplied text (phones like "+1 610 555 0123" stay intact).
    starts_formula = text[:1] in ("=", "@") or (text[:1] in ("+", "-") and any(c.isalpha() for c in text))
    return "'" + text if starts_formula else text


class ExportService:
    def __init__(self):
        self._lock = threading.Lock()
        self._last_snapshot = 0.0

    @property
    def event_dir(self) -> str:
        os.makedirs(settings.storage.event_dir, exist_ok=True)
        return settings.storage.event_dir

    @property
    def backups_dir(self) -> str:
        path = os.path.join(self.event_dir, "backups")
        os.makedirs(path, exist_ok=True)
        return path

    # ------------------------------------------------------------------ helpers
    def _atomic_write(self, final_path: str, writer):
        fd, tmp = tempfile.mkstemp(dir=os.path.dirname(final_path), prefix=".tmp_", suffix=os.path.splitext(final_path)[1])
        os.close(fd)
        try:
            writer(tmp)
            os.replace(tmp, final_path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)

    def _rows(self, records: List[Dict[str, Any]], columns) -> List[List[Any]]:
        return [[_cell(r.get(key)) for _, key in columns] for r in records]

    def _collect(self):
        from app.services.db_service import db_service
        perfs = db_service.get_all_performances()
        food = db_service.get_all_food_items_with_signups()
        for f in food:
            f["_status"] = "Taken" if f.get("is_taken") else "Open"
        try:
            participants = db_service.participant_repo.get_all_participants(all_performances=[], all_food_items=[])
        except Exception:
            participants = []
        return perfs, food, participants

    # ------------------------------------------------------------------ outputs
    def _write_xlsx(self, path: str, perfs, food, participants):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="C2410C")

        def add_sheet(title, columns, records, first=False):
            ws = wb.active if first else wb.create_sheet()
            ws.title = title
            ws.append([c for c, _ in columns])
            for row in self._rows(records, columns):
                ws.append(row)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            for idx, (header, _) in enumerate(columns, start=1):
                longest = max([len(str(header))] + [len(str(r[idx - 1])) for r in self._rows(records, columns)] or [10])
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 10), 48)
            return ws

        add_sheet("Performances", PERFORMANCE_COLUMNS, perfs, first=True)
        run_sheet_cols = [("Seq", "sequence_order"), ("Performer(s)", "_who"), ("Song", "song_title"),
                          ("Type", "performance_type"), ("Track", "track_status"), ("Notes", "stage_notes")]
        run = []
        for p in sorted(perfs, key=lambda x: x.get("sequence_order") or 9999):
            who = p.get("performer_name") or ""
            if p.get("partner_name"):
                who += f" & {p['partner_name']}"
            run.append({**p, "_who": who})
        add_sheet("Run Sheet", run_sheet_cols, run)
        add_sheet("Potluck", POTLUCK_COLUMNS, food)
        add_sheet("Participants", PARTICIPANT_COLUMNS, participants)
        wb.save(path)

    def _write_csv(self, path: str, columns, records):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([c for c, _ in columns])
            w.writerows(self._rows(records, columns))

    def snapshot_database(self) -> str:
        from app.services.db_service import db_service
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        dest = os.path.join(self.backups_dir, f"event_{stamp}.db")
        src = db_service._get_connection()
        dst = sqlite3.connect(dest)
        try:
            src.backup(dst)
        finally:
            dst.close()
        keep = max(int(getattr(settings.backup, "keep_snapshots", 30)), 1)
        for old in sorted(glob.glob(os.path.join(self.backups_dir, "event_*.db")))[:-keep]:
            try:
                os.remove(old)
            except OSError:
                pass
        self._last_snapshot = datetime.now().timestamp()
        return dest

    # ------------------------------------------------------------------ public
    def export_now(self, snapshot: bool = True, force_snapshot: bool = False) -> Dict[str, Any]:
        with self._lock:
            perfs, food, participants = self._collect()
            base = self.event_dir
            self._atomic_write(os.path.join(base, "signups.xlsx"), lambda p: self._write_xlsx(p, perfs, food, participants))
            self._atomic_write(os.path.join(base, "signups.csv"), lambda p: self._write_csv(p, PERFORMANCE_COLUMNS, perfs))
            self._atomic_write(os.path.join(base, "potluck.csv"), lambda p: self._write_csv(p, POTLUCK_COLUMNS, food))
            snap = None
            now = datetime.now().timestamp()
            if snapshot and (force_snapshot or now - self._last_snapshot >= SNAPSHOT_MIN_INTERVAL_SECONDS):
                snap = self.snapshot_database()
            return {
                "rows_backed": len(perfs) + len(food),
                "performances": len(perfs),
                "potluck_items": len(food),
                "files": ["signups.xlsx", "signups.csv", "potluck.csv"],
                "snapshot": os.path.basename(snap) if snap else None,
                "event_dir": base,
            }


export_service = ExportService()
